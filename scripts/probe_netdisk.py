# -*- coding: utf-8 -*-
"""只读探测：下载资料库网盘《账号名称.xlsx》兜底文件，检查「工作表2」最新日期与行数。"""
import io, os, sys, json, urllib.request, runpy

LIB = r"C:\Users\张彦飞\AppData\Local\Programs\WorkBuddy\resources\app.asar.unpacked\resources\plugins\workbuddy-builtin\skills\library"
BASE = r"D:\workbuddy\2026-09-18-22-20-37\delivery"
XLSX_NODE = "FbaXm0rOJGeHRgCAIwVM4G"

token = (sys.stdin.readline() or "").strip()


def run_script(rel, args):
    script = os.path.join(LIB, rel)
    old_argv, old_stdin, old_stdout = sys.argv, sys.stdin, sys.stdout
    sys.argv = [script] + args
    sys.stdin = io.StringIO(token + "\n")
    buf = io.StringIO()
    sys.stdout = buf
    try:
        runpy.run_path(script, run_name="__main__")
    except SystemExit:
        pass
    except Exception as e:
        buf.write("[EXC] %r" % (e,))
    finally:
        sys.argv, sys.stdin, sys.stdout = old_argv, old_stdin, old_stdout
    return buf.getvalue()


out = run_script(os.path.join("drive", "get_download_link.py"),
                 ["--token-stdin", "--node-id", XLSX_NODE])
dl_url = None
for ln in out.splitlines():
    if ln.startswith("KS_DRIVE_DOWNLOAD"):
        dl_url = json.loads(ln.split("KS_DRIVE_DOWNLOAD", 1)[1].strip())["download_url"]
        break
if not dl_url:
    print("PROBE_FAIL: no download url -> " + out[:300])
    sys.exit(0)

import openpyxl
tmp = os.path.join(BASE, "data", "_probe_tmp.xlsx")
req = urllib.request.Request(dl_url, headers={"User-Agent": "library-sync/1.0"})
with open(tmp, "wb") as f:
    f.write(urllib.request.urlopen(req, timeout=120).read())
wb = openpyxl.load_workbook(tmp, data_only=True)
print("sheets:", wb.sheetnames)
if "工作表2" not in wb.sheetnames:
    print("PROBE_FAIL: no 工作表2")
    sys.exit(0)
rows = list(wb["工作表2"].iter_rows(values_only=True))
from datetime import datetime, timedelta
dates = []
for r in rows[1:]:
    if not r or not r[0] or len(r) < 2:
        continue
    v = r[1]
    if hasattr(v, "strftime"):
        dates.append(v.strftime("%Y-%m-%d"))
    elif isinstance(v, (int, float)):
        dates.append((datetime(1899, 12, 30) + timedelta(days=float(v))).strftime("%Y-%m-%d"))
    else:
        dates.append(str(v)[:10])
print("PROBE_OK raw_rows=%d valid_rows=%d max_date=%s" % (len(rows), len(dates), max(dates) if dates else "none"))
