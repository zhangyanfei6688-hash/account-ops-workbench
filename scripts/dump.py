# -*- coding: utf-8 -*-
import openpyxl

p = r"D:\workbuddy\2026-09-18-22-20-37\delivery\data\account.xlsx"
out = r"D:\workbuddy\2026-09-18-22-20-37\delivery\data\dump.txt"
wb = openpyxl.load_workbook(p, data_only=True)
lines = []
lines.append("SHEETS: " + repr(wb.sheetnames))
for ws in wb.worksheets:
    lines.append("=== SHEET: %s | dims=%s | max_row=%s | max_col=%s" % (ws.title, ws.dimensions, ws.max_row, ws.max_column))
    n = 0
    for row in ws.iter_rows(values_only=True):
        cells = ["" if c is None else str(c) for c in row]
        lines.append("R%d: %s" % (n + 1, " | ".join(cells)))
        n += 1
        if n > 400:
            lines.append("...TRUNCATED...")
            break
with open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("done")
